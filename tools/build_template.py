"""
Generate src/OrthoLink.AddIn/template.pptx.

The template holds one "custom geometry" shape per connector topology
(sequence of horizontal / vertical segments). The add-in copies the right
shape into the user's slide, then only touches position / size / flip /
adjust values, never the geometry itself.

Adjust values (visible in the object model as Shape.Adjustments):
  adj1          corner radius in EMU (12700 EMU = 1 pt); no handle
  adj2 .. adjN  position of each *interior* segment as a fraction of the
                shape width (vertical segments) or height (horizontal
                segments); each has a yellow drag handle in PowerPoint

Local coordinate space of every shape: the path runs from (0,0) to (w,h).
PowerPoint's flipH / flipV take care of the other three quadrants.
Every corner is a quarter circle whose radius shrinks automatically when a
segment is too short.
"""
import os, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
DEFAULT_RADIUS_EMU = int(Pt(12))
Q = 5400000          # 90 degrees in 60000ths of a degree
TOPOLOGIES = ["HV", "VH", "HVH", "VHV", "HVHV", "VHVH", "HVHVH", "VHVHV"]


def gd(name, fmla):
    return f'<a:gd name="{name}" fmla="{fmla}"/>'


def build(topo):
    """Return (guides, handles, path) for a segment sequence like 'HVHV'."""
    segs = list(topo)
    n = len(segs)
    guides = [gd("rad", "*/ adj1 1 1")]
    handles = []

    # position of each segment along its perpendicular axis (H: y, V: x)
    pos = []
    free = 0
    for k, s in enumerate(segs):
        if k == 0:
            pos.append("0")
        elif k == n - 1:
            pos.append("h" if s == "H" else "w")
        else:
            free += 1
            name = f"p{k}"
            guides.append(gd(name, f"*/ {'h' if s == 'H' else 'w'} adj{free + 1} 100000"))
            pos.append(name)

    # corner points C1..C(n-1): corner i sits between seg i-1 and seg i
    cx, cy = ["0"], ["0"]
    for i in range(1, n):
        a, b = segs[i - 1], segs[i]
        if a == "H":     # H then V: x from the V segment, y from the H segment
            cx.append(pos[i]); cy.append(pos[i - 1])
        else:
            cx.append(pos[i - 1]); cy.append(pos[i])
    cx.append("w"); cy.append("h")

    # segment deltas, directions, lengths
    lens = []
    dirs = []
    for k, s in enumerate(segs):
        if s == "H":
            guides.append(gd(f"d{k}", f"+- {cx[k + 1]} 0 {cx[k]}"))
        else:
            guides.append(gd(f"d{k}", f"+- {cy[k + 1]} 0 {cy[k]}"))
        guides.append(gd(f"len{k}", f"abs d{k}"))
        guides.append(gd(f"dir{k}", f"?: d{k} 1 -1"))
        lens.append(f"len{k}"); dirs.append(f"dir{k}")

    # radius: min(rad, terminal segment lengths, half of interior segment lengths)
    prev = "rad"
    for k in range(n):
        if 0 < k < n - 1:
            guides.append(gd(f"half{k}", f"*/ {lens[k]} 1 2"))
            lim = f"half{k}"
        else:
            lim = lens[k]
        guides.append(gd(f"rm{k}", f"min {prev} {lim}"))
        prev = f"rm{k}"
    guides.append(gd("rEff", f"max {prev} 1"))

    # path
    path = ['<a:moveTo><a:pt x="l" y="t"/></a:moveTo>']
    for i in range(1, n):
        din, dout = dirs[i - 1], dirs[i]
        guides.append(gd(f"off{i}", f"*/ rEff {din} 1"))
        if segs[i - 1] == "H":           # H -> V corner
            guides.append(gd(f"tx{i}", f"+- {cx[i]} 0 off{i}"))
            path.append(f'<a:lnTo><a:pt x="tx{i}" y="{cy[i]}"/></a:lnTo>')
            guides.append(gd(f"st{i}", f"?: {dout} 16200000 {Q}"))
            guides.append(gd(f"pr{i}", f"*/ {din} {dout} 1"))
            guides.append(gd(f"sw{i}", f"*/ {Q} pr{i} 1"))
        else:                            # V -> H corner
            guides.append(gd(f"ty{i}", f"+- {cy[i]} 0 off{i}"))
            path.append(f'<a:lnTo><a:pt x="{cx[i]}" y="ty{i}"/></a:lnTo>')
            guides.append(gd(f"st{i}", f"?: {dout} 10800000 0"))
            guides.append(gd(f"pr{i}", f"*/ {din} {dout} 1"))
            guides.append(gd(f"sw{i}", f"*/ -{Q} pr{i} 1"))
        path.append(f'<a:arcTo wR="rEff" hR="rEff" stAng="st{i}" swAng="sw{i}"/>')
    path.append('<a:lnTo><a:pt x="r" y="b"/></a:lnTo>')

    # handles: one per interior segment, sitting at the middle of that segment
    free = 0
    for k in range(1, n - 1):
        free += 1
        if segs[k] == "V":
            guides.append(gd(f"hm{k}", f"+/ {cy[k]} {cy[k + 1]} 2"))
            handles.append(f'<a:ahXY gdRefX="adj{free + 1}" minX="-2147483647" maxX="2147483647">'
                           f'<a:pos x="{pos[k]}" y="hm{k}"/></a:ahXY>')
        else:
            guides.append(gd(f"hm{k}", f"+/ {cx[k]} {cx[k + 1]} 2"))
            handles.append(f'<a:ahXY gdRefY="adj{free + 1}" minY="-2147483647" maxY="2147483647">'
                           f'<a:pos x="hm{k}" y="{pos[k]}"/></a:ahXY>')
    return guides, handles, "".join(path), free


def custgeom_xml(topo):
    guides, handles, path, free = build(topo)
    av = [gd("adj1", f"val {DEFAULT_RADIUS_EMU}")] + [gd(f"adj{i + 2}", "val 50000") for i in range(free)]
    return (
        f'<a:custGeom xmlns:a="{A}">'
        f'<a:avLst>{"".join(av)}</a:avLst>'
        f'<a:gdLst>{"".join(guides)}</a:gdLst>'
        f'<a:ahLst>{"".join(handles)}</a:ahLst>'
        '<a:cxnLst/>'
        '<a:rect l="l" t="t" r="r" b="b"/>'
        f'<a:pathLst><a:path fill="none">{path}</a:path></a:pathLst>'
        '</a:custGeom>'
    )


def main(out_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    x = Pt(20)
    for topo in TOPOLOGIES:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Pt(60), Pt(100), Pt(80))
        shp.name = "OL_" + topo
        spPr = shp._element.spPr
        prst = spPr.find(qn("a:prstGeom"))
        prst.addprevious(etree.fromstring(custgeom_xml(topo)))
        spPr.remove(prst)
        shp.fill.background()
        shp.line.width = Pt(1.5)
        shp.line.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0x40, 0x40, 0x40)
        ln = spPr.find(qn("a:ln"))
        etree.SubElement(ln, qn("a:round"))
        for tag in ("p:txBody", "p:style"):
            el = shp._element.find(qn(tag))
            if el is not None:
                shp._element.remove(el)
        x += Pt(115)
    prs.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    default = os.path.join(here, "..", "src", "OrthoLink.AddIn", "template.pptx")
    main(os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else default))
